"""Build a formatted Excel workbook from the counterbalanced design (v2 – 2 scenarios)."""

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from generate_design import HIGH_MSG, LOW_MSG, build

OUT  = os.path.dirname(os.path.abspath(__file__))
PATH = os.path.join(OUT, "study_design.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")

# Interview = orange, Mathe = blue
SCEN_FILL = {
    "Interview": PatternFill("solid", fgColor="FCE4D6"),   # orange
    "Mathe":     PatternFill("solid", fgColor="DDEBF7"),   # blue
}

# System accent colours (subtle tint for the Trials sheet)
SYS_FONT = {
    "Baseline-Earcon":  Font(color="375623"),   # dark green
    "Baseline-Speech":  Font(color="1F3864"),   # dark blue
    "SonoAdapt":        Font(bold=True, color="7B2C2C"),  # dark red
}

THIN      = Side(style="thin",   color="BFBFBF")
THICK_TOP = Side(style="medium", color="505050")
BORDER    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_BT = Border(left=THIN, right=THIN,
                   top=THICK_TOP, bottom=THIN)    # new participant marker


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------------------
# Sheet 1 – Lookup  (1 row per participant, columns B01–B12)
# ---------------------------------------------------------------------------

def sheet_lookup(wb, rows):
    ws = wb.create_sheet("Lookup")
    ws.append(["Participant"] + [f"B{i:02d}" for i in range(1, 13)])
    style_header(ws)

    by_p = {}
    for r in rows:
        by_p.setdefault(r["participant"], {})[r["block"]] = r

    for p in sorted(by_p):
        ws.append([p] + [by_p[p][i]["code"] for i in range(1, 13)])
        excel_row = ws.max_row
        for col in range(1, 14):
            c = ws.cell(row=excel_row, column=col)
            c.alignment = Alignment(horizontal="center")
            c.border    = BORDER
            c.font      = Font(name="Consolas", size=11)
            if col > 1:
                trial_data = by_p[p][col - 1]
                c.fill = SCEN_FILL[trial_data["scenario"]]
            else:
                c.font = Font(bold=True)

    set_widths(ws, [12] + [8] * 12)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:M{ws.max_row}"

    note = ws.max_row + 2
    ws.cell(row=note, column=1,
            value="Code = <Scenario I|M><System E|S|X><Urgency L|H><Message 1-6>   "
                  "e.g. MSH4 = Mathe / Baseline-Speech / HIGH / H4 (Paul)").font = Font(italic=True)
    ws.cell(row=note + 1, column=1,
            value="Cell colour = scenario:  orange = Interview,  blue = Mathe").font = Font(italic=True)
    ws.cell(row=note + 2, column=1,
            value="Blocks 1-6 = first scenario block, Blocks 7-12 = second scenario block").font = Font(italic=True)
    return ws


# ---------------------------------------------------------------------------
# Sheet 2 – Trials  (1 row per trial, full detail)
# ---------------------------------------------------------------------------

def sheet_trials(wb, rows):
    ws = wb.create_sheet("Trials")
    headers = [
        "Participant", "Block", "Code", "Scenario block", "Trial in block",
        "Scenario", "System", "Urgency", "Msg", "Message text",
    ]
    ws.append(headers)
    style_header(ws)

    prev_p = None
    for r in rows:
        ws.append([
            r["participant"], r["block"], r["code"],
            r["scenario_block"], r["trial_in_scenario"],
            r["scenario"], r["system"], r["urgency"],
            r["message_id"], r["message_text"],
        ])
        excel_row = ws.max_row
        new_p     = r["participant"] != prev_p
        prev_p    = r["participant"]

        for col in range(1, len(headers) + 1):
            c = ws.cell(row=excel_row, column=col)
            c.border    = BORDER_BT if new_p else BORDER
            if col != 10:
                c.alignment = Alignment(horizontal="center")

        # Scenario colour band
        ws.cell(row=excel_row, column=6).fill = SCEN_FILL[r["scenario"]]
        # Code cell – monospace bold
        ws.cell(row=excel_row, column=3).font = Font(name="Consolas", size=11, bold=True)
        # HIGH urgency highlighted
        if r["urgency"] == "HIGH":
            ws.cell(row=excel_row, column=8).font = Font(bold=True, color="C00000")
        # System colour hint
        ws.cell(row=excel_row, column=7).font = SYS_FONT.get(r["system"], Font())

    set_widths(ws, [11, 7, 8, 14, 14, 11, 18, 10, 7, 95])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{ws.max_row}"
    return ws


# ---------------------------------------------------------------------------
# Sheet 3 – Legend
# ---------------------------------------------------------------------------

def sheet_legend(wb):
    ws = wb.create_sheet("Legend")
    ws.append(["Element", "Code", "Meaning"])
    style_header(ws)

    data = [
        ("Scenario", "I", "Interview"),
        ("Scenario", "M", "Mathe / Rechen-Task"),
        ("System",   "E", "Baseline-Earcon"),
        ("System",   "S", "Baseline-Speech"),
        ("System",   "X", "SonoAdapt"),
        ("Urgency",  "L", "LOW"),
        ("Urgency",  "H", "HIGH"),
    ]
    for row in data:
        ws.append(row)

    ws.append([])
    ws.append(["LOW message", "L1..L6", "text"])
    style_header(ws, ws.max_row)
    for i, txt in LOW_MSG.items():
        ws.append(["LOW", f"L{i}", txt])

    ws.append([])
    ws.append(["HIGH message", "H1..H6", "text"])
    style_header(ws, ws.max_row)
    for i, txt in HIGH_MSG.items():
        ws.append(["HIGH", f"H{i}", txt])

    ws.append([])
    ws.append(["Qualtrics",   "RegEx", "^[IM][ESX][LH][1-6]$"])
    ws.append(["Sample size", "N",     "use multiples of 6 (6 / 12 / 24) for full balance"])
    ws.append(["Block structure", "–", "Trials 1-6 = first scenario block, 7-12 = second"])

    for row in ws.iter_rows(min_row=2, max_col=3):
        for c in row:
            if c.column == 2:
                c.font      = Font(name="Consolas", size=11)
                c.alignment = Alignment(horizontal="center")

    set_widths(ws, [16, 12, 95])
    return ws


# ---------------------------------------------------------------------------
# Sheet 4 – Per-participant run card  (one mini-table per participant)
#            Printed version for the study moderator.
# ---------------------------------------------------------------------------

def sheet_runcards(wb, rows):
    """Compact per-participant summary – one block of rows per participant."""
    ws = wb.create_sheet("Run Cards")

    by_p = {}
    for r in rows:
        by_p.setdefault(r["participant"], []).append(r)

    for p in sorted(by_p):
        # Participant header
        ws.append([f"Participant {p:02d}"])
        hdr_row = ws.max_row
        ws.cell(hdr_row, 1).font      = Font(bold=True, size=13, color="FFFFFF")
        ws.cell(hdr_row, 1).fill      = HEADER_FILL
        ws.cell(hdr_row, 1).alignment = Alignment(horizontal="left")
        ws.merge_cells(start_row=hdr_row, start_column=1,
                       end_row=hdr_row, end_column=5)

        # Column headers
        ws.append(["Block", "Code", "Scenario", "System", "Urgency", "Msg", "Message text"])
        col_hdr = ws.max_row
        for c in ws[col_hdr]:
            c.font      = Font(bold=True, color="FFFFFF")
            c.fill      = PatternFill("solid", fgColor="4472C4")
            c.alignment = Alignment(horizontal="center")
            c.border    = BORDER

        for r in by_p[p]:
            ws.append([
                r["block"], r["code"], r["scenario"], r["system"],
                r["urgency"], r["message_id"], r["message_text"],
            ])
            data_row = ws.max_row
            for col in range(1, 8):
                c = ws.cell(data_row, col)
                c.border    = BORDER
                c.alignment = Alignment(horizontal="center" if col < 7 else "left")
            ws.cell(data_row, 2).font = Font(name="Consolas", size=11, bold=True)
            ws.cell(data_row, 3).fill = SCEN_FILL[r["scenario"]]
            if r["urgency"] == "HIGH":
                ws.cell(data_row, 5).font = Font(bold=True, color="C00000")

        ws.append([])   # blank row between participants

    set_widths(ws, [7, 9, 12, 20, 10, 7, 90])
    ws.freeze_panes = "A1"
    return ws


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    rows = build()
    wb   = Workbook()
    wb.remove(wb.active)
    sheet_lookup(wb, rows)
    sheet_trials(wb, rows)
    sheet_runcards(wb, rows)
    sheet_legend(wb)
    wb.active = 0
    wb.save(PATH)
    print("wrote", PATH)


if __name__ == "__main__":
    main()
